import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parent))
import import_lt_table_xlsx as importer


class ImportLtTableXlsxTest(unittest.TestCase):
    def test_read_cli_connection_prefers_environment_host_and_token(self):
        with patch.dict(
            os.environ,
            {
                "LINGTONG_HOST": "http://test.example:8085/",
                "LINGTONG_TOKEN": "apk-test",
            },
            clear=False,
        ):
            self.assertEqual(
                ("http://test.example:8085", "apk-test"),
                importer.read_cli_connection(),
            )

    def test_workflow_script_does_not_use_blocked_java_classes(self):
        workflow_path = Path(__file__).parent / "lt_table_batch_import_workflow.json"
        workflow = json.loads(workflow_path.read_text(encoding="utf-8"))
        script = workflow["nodes"][1]["data"]["scriptConfig"]["script"]

        self.assertNotIn("java.util", script)
        self.assertIn("AppInvoker.invoke", script)

    def test_normalize_identifier_preserves_large_integer(self):
        self.assertEqual(
            "6955294126991677326", importer.normalize_text(6955294126991677326)
        )
        self.assertEqual("5982188428147034", importer.normalize_text(5982188428147034.0))

    def test_normalize_date_uses_shanghai_timezone(self):
        expected = int(
            datetime(2026, 9, 13, 23, 59, 52, tzinfo=importer.SHANGHAI).timestamp()
            * 1000
        )
        self.assertEqual(expected, importer.normalize_date("2026-09-13 23:59:52"))
        self.assertIsNone(importer.normalize_date(""))

    def test_build_record_maps_schema_and_import_key(self):
        field_ids = {
            title: index + 2 for index, title in enumerate(importer.SOURCE_HEADERS)
        }
        field_ids[importer.IMPORT_KEY_TITLE] = 1
        row = [None] * len(importer.SOURCE_HEADERS)
        row[0] = 7.0
        row[1] = 6018232734063081
        row[8] = "2026-09-13 23:59:52"
        row[17] = 2.0

        record = importer.build_record(
            row,
            field_ids=field_ids,
            schema_id=99,
            schema_version=3,
            app_id=213,
            import_batch="20260914222134",
        )

        self.assertEqual("20260914222134:7", record["primaryValue"])
        self.assertEqual("20260914222134:7", record["data"]["1"])
        self.assertEqual("6018232734063081", record["data"]["3"])
        self.assertEqual(2, record["data"]["19"])
        self.assertIsInstance(record["data"]["10"], int)

    def test_build_record_preserves_multiple_aftersale_timestamps_as_text(self):
        field_ids = {
            title: index + 2 for index, title in enumerate(importer.SOURCE_HEADERS)
        }
        field_ids[importer.IMPORT_KEY_TITLE] = 1
        row = [None] * len(importer.SOURCE_HEADERS)
        row[0] = 1501
        row[24] = "2026-09-13 21:25:56,2026-09-13 21:25:59"

        record = importer.build_record(
            row,
            field_ids=field_ids,
            schema_id=99,
            schema_version=3,
            app_id=213,
            import_batch="20260914222134",
        )

        self.assertEqual(
            "2026-09-13 21:25:56,2026-09-13 21:25:59",
            record["data"]["26"],
        )

    def test_schema_has_hidden_import_key_and_no_platform_primary_key(self):
        columns = importer.build_columns_schema()

        expected_count = (
            len(importer.SOURCE_HEADERS)
            + len(importer.PRODUCT_HEADERS)
            - len(importer.PRODUCT_EXCLUDED_HEADERS)
            + 1
        )
        self.assertEqual(expected_count, len(columns))
        self.assertEqual(importer.IMPORT_KEY_TITLE, columns[0]["title"])
        self.assertFalse(columns[0]["primaryKey"])
        self.assertFalse(columns[0]["show"])
        self.assertFalse(any(column["primaryKey"] for column in columns))
        self.assertEqual(len(columns), len({column["key"] for column in columns}))
        titles = {column["title"] for column in columns}
        self.assertIn("商品名称(档案)", titles)
        self.assertIn("商品品牌(档案)", titles)
        self.assertNotIn("序号(档案)", titles)
        self.assertNotIn("SKU聚合键(档案)", titles)
        self.assertNotIn("规格商家编码(档案)", titles)

    def test_load_product_lookup_keeps_requested_skus_only(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "products.xlsx"
            self._write_product_workbook(path)

            lookup = importer.load_product_lookup(path, {"SKU-2"})

        self.assertEqual({"SKU-2"}, set(lookup))
        self.assertEqual("商品二", lookup["SKU-2"]["商品名称"])
        self.assertEqual(2026, lookup["SKU-2"]["年份"])

    def test_load_product_lookup_rejects_missing_order_sku(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "products.xlsx"
            self._write_product_workbook(path)

            with self.assertRaisesRegex(importer.ImportFailure, "SKU-3"):
                importer.load_product_lookup(path, {"SKU-1", "SKU-3"})

    def test_build_record_appends_product_archive_fields(self):
        columns = importer.build_columns_schema()
        field_ids = {column["title"]: column["id"] for column in columns}
        row = [None] * len(importer.SOURCE_HEADERS)
        row[0] = 7
        row[5] = "SKU-1"
        product = {header: None for header in importer.PRODUCT_HEADERS}
        product.update(
            {
                "SKU聚合键": "SKU-1",
                "商品名称": "档案商品名",
                "规格成本价": 12.5,
                "业务上市日期": "2026-08-01",
                "年份": 2026,
            }
        )

        record = importer.build_record(
            row,
            product=product,
            field_ids=field_ids,
            schema_id=99,
            schema_version=3,
            app_id=213,
            import_batch="20260914222134-product",
        )

        self.assertEqual(
            "档案商品名", record["data"][str(field_ids["商品名称(档案)"])]
        )
        self.assertEqual(
            12.5, record["data"][str(field_ids["规格成本价(档案)"])]
        )
        self.assertIsInstance(
            record["data"][str(field_ids["业务上市日期(档案)"])], int
        )
        self.assertEqual(2026, record["data"][str(field_ids["年份(档案)"])])

    def test_iter_source_rows_orders_files_and_skips_summary(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            base = Path(temp_dir) / "report"
            self._write_workbook(base.with_suffix(".xlsx"), 1, 2)
            self._write_workbook(Path(f"{base}-1.xlsx"), 3, 4)
            self._write_workbook(Path(f"{base}-2.xlsx"), 5, 5, include_summary=True)

            rows = list(importer.iter_source_rows(Path(temp_dir)))

        self.assertEqual([1, 2, 3, 4, 5], [int(row[0]) for row in rows])

    def test_validate_workflow_response_rejects_connector_failure(self):
        with self.assertRaisesRegex(importer.ImportFailure, "connector rejected"):
            importer.validate_workflow_response(
                {
                    "success": True,
                    "result": {
                        "result": {"code": 500, "msg": "connector rejected"}
                    },
                }
            )

    def test_build_workflow_request_uses_connector_account_environment(self):
        payload = importer.build_workflow_request(
            account_name="纤莉秀表格导入",
            table_id=1747,
            records=[],
        )

        self.assertEqual("prod", payload["env"])

    def test_classify_failed_batch_count_distinguishes_retry_and_commit(self):
        self.assertEqual(
            "retry", importer.classify_failed_batch_count(1000, 500, 1000)
        )
        self.assertEqual(
            "committed", importer.classify_failed_batch_count(1000, 500, 1500)
        )
        with self.assertRaisesRegex(importer.ImportFailure, "部分提交"):
            importer.classify_failed_batch_count(1000, 500, 1200)

    def test_reconcile_remote_progress_recovers_committed_rows(self):
        self.assertEqual(
            (1100510, 1100510),
            importer.reconcile_remote_progress(
                processed=1100010,
                last_sequence=1100010,
                remote_count=1100510,
                last_import_key="20260915-product-archive:1100510",
                import_batch="20260915-product-archive",
            ),
        )

    def test_reconcile_remote_progress_rejects_unexpected_import_key(self):
        with self.assertRaisesRegex(importer.ImportFailure, "追踪键"):
            importer.reconcile_remote_progress(
                processed=1100010,
                last_sequence=1100010,
                remote_count=1100510,
                last_import_key="another-batch:1100510",
                import_batch="20260915-product-archive",
            )

    @staticmethod
    def _write_workbook(path: Path, start: int, end: int, include_summary: bool = False):
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet()
        sheet.append(["report"] * len(importer.SOURCE_HEADERS))
        sheet.append(importer.SOURCE_HEADERS)
        for sequence in range(start, end + 1):
            row = [None] * len(importer.SOURCE_HEADERS)
            row[0] = sequence
            row[1] = 1000 + sequence
            sheet.append(row)
        if include_summary:
            sheet.append(["汇总"] + [None] * (len(importer.SOURCE_HEADERS) - 1))
        workbook.save(path)

    @staticmethod
    def _write_product_workbook(path: Path):
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("商品档案")
        sheet.append(importer.PRODUCT_HEADERS)
        for index in range(1, 3):
            row = [None] * len(importer.PRODUCT_HEADERS)
            row[0] = index
            row[1] = f"SKU-{index}"
            row[2] = f"MAIN-{index}"
            row[3] = f"商品{['一', '二'][index - 1]}"
            row[5] = f"SKU-{index}"
            row[47] = 2026
            sheet.append(row)
        workbook.save(path)


if __name__ == "__main__":
    unittest.main()
