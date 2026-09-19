#!/usr/bin/env python3
"""Stream large order-detail workbooks into a Lingtong table workflow."""

from __future__ import annotations

import argparse
import base64
import hashlib
import json
import math
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request
import warnings
from datetime import date, datetime, time as datetime_time
from pathlib import Path
from typing import Any, Iterator
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

warnings.filterwarnings("ignore", message="Workbook contains no default style")


SOURCE_HEADERS = [
    "序号",
    "订单商品明细ID",
    "系统订单号",
    "平台订单号",
    "商品图片(订单)",
    "系统商品编码(订单)",
    "系统商品标题(订单)",
    "系统规格属性名称(订单)",
    "付款时间",
    "发货时间",
    "完成时间",
    "店铺平台",
    "店铺名称",
    "分销商",
    "分销商简称",
    "订单状态(明细)",
    "平台状态",
    "订单商品数",
    "系统实发商品数",
    "订单买家已付金额",
    "系统实发金额",
    "销售成本",
    "系统实发成本",
    "售后工单号",
    "申请时间(售后)",
    "完成时间(售后)",
    "售后类型",
    "申请数量",
    "退货数量",
    "实退数量",
    "退货数量(排除仅退款)",
    "退款金额",
    "退款成本",
    "退货成本(排除仅退款)",
    "实退成本",
]

SOURCE_KEYS = [
    "sequence_no",
    "order_item_detail_id",
    "system_order_no",
    "platform_order_no",
    "order_image_url",
    "system_product_code",
    "system_product_title",
    "system_sku_name",
    "paid_at",
    "shipped_at",
    "completed_at",
    "shop_platform",
    "shop_name",
    "distributor",
    "distributor_short_name",
    "order_item_status",
    "platform_status",
    "ordered_qty",
    "shipped_qty",
    "buyer_paid_amount",
    "shipped_amount",
    "sales_cost",
    "shipped_cost",
    "aftersale_ticket_no",
    "aftersale_applied_at",
    "aftersale_completed_at",
    "aftersale_type",
    "aftersale_applied_qty",
    "return_qty",
    "actual_return_qty",
    "return_qty_excl_refund_only",
    "refund_amount",
    "refund_cost",
    "return_cost_excl_refund_only",
    "actual_return_cost",
]

PRODUCT_HEADERS = [
    "序号",
    "SKU聚合键",
    "主商家编码",
    "商品名称",
    "是否含SKU",
    "规格商家编码",
    "SKU图片",
    "商品图片",
    "颜色属性",
    "第二属性",
    "第三属性",
    "第四属性",
    "规格成本价",
    "规格销售价",
    "规格批发价",
    "规格市场价",
    "规格别名",
    "规格简称",
    "条形码",
    "规格备注",
    "供应商名称",
    "供应商编码",
    "供应商进价",
    "基本单位",
    "重量(kg)",
    "长(cm)",
    "宽(cm)",
    "高(cm)",
    "体积(cm³)",
    "装箱数",
    "负责人",
    "商品简称",
    "SKU一级分类",
    "SKU二级分类",
    "SKU三级分类",
    "商品一级分类",
    "商品二级分类",
    "商品三级分类",
    "商品根分类",
    "商品二级类目",
    "商品三级类目",
    "商品四级类目",
    "SKU品牌",
    "商品品牌",
    "营销波段",
    "业务上市日期",
    "季节",
    "年份",
    "设计师",
    "ERP上架日期",
    "商品创建时间",
    "SKU创建时间",
    "商品类型",
]

PRODUCT_EXCLUDED_HEADERS = {"序号", "SKU聚合键", "规格商家编码"}
PRODUCT_DATE_HEADERS = {
    "业务上市日期",
    "ERP上架日期",
    "商品创建时间",
    "SKU创建时间",
}
PRODUCT_NUMBER_HEADERS = {
    "规格成本价",
    "规格销售价",
    "规格批发价",
    "规格市场价",
    "供应商进价",
    "重量(kg)",
    "长(cm)",
    "宽(cm)",
    "高(cm)",
    "体积(cm³)",
    "装箱数",
    "年份",
}

IMPORT_KEY_TITLE = "导入唯一键"
DATE_HEADERS = {
    "付款时间",
    "发货时间",
    "完成时间",
}
NUMBER_HEADERS = {
    "序号",
    "订单商品数",
    "系统实发商品数",
    "订单买家已付金额",
    "系统实发金额",
    "销售成本",
    "系统实发成本",
    "申请数量",
    "退货数量",
    "实退数量",
    "退货数量(排除仅退款)",
    "退款金额",
    "退款成本",
    "退货成本(排除仅退款)",
    "实退成本",
}
SHANGHAI = ZoneInfo("Asia/Shanghai")
KEYRING_BASE64_PREFIX = "go-keyring-base64:"
KEYRING_HEX_PREFIX = "go-keyring-encoded:"


class ImportFailure(RuntimeError):
    """Raised when source validation or a remote import batch fails."""


def build_columns_schema() -> list[dict[str, Any]]:
    columns = [
        {
            "id": 1,
            "title": IMPORT_KEY_TITLE,
            "key": "import_key",
            "type": "text",
            "primaryKey": False,
            "show": False,
            "width": 220,
            "align": "left",
            "description": "导入批次与原始序号组成的追踪键",
        }
    ]
    amount_headers = {
        "订单买家已付金额",
        "系统实发金额",
        "销售成本",
        "系统实发成本",
        "退款金额",
        "退款成本",
        "退货成本(排除仅退款)",
        "实退成本",
    }
    for index, (title, key) in enumerate(zip(SOURCE_HEADERS, SOURCE_KEYS), start=2):
        column: dict[str, Any] = {
            "id": index,
            "title": title,
            "key": key,
            "primaryKey": False,
            "show": True,
            "width": 180,
            "align": "left",
        }
        if title in DATE_HEADERS:
            column.update(type="date", dateFormat="YYYY-MM-DD HH:mm:ss")
        elif title in NUMBER_HEADERS:
            number_format = "1,000.00" if title in amount_headers else "1000"
            column.update(type="number", numberFormat=number_format, align="right")
        else:
            column["type"] = "text"
        columns.append(column)

    next_id = len(columns) + 1
    product_amount_headers = {
        "规格成本价",
        "规格销售价",
        "规格批发价",
        "规格市场价",
        "供应商进价",
    }
    for index, title in enumerate(PRODUCT_HEADERS):
        if title in PRODUCT_EXCLUDED_HEADERS:
            continue
        column = {
            "id": next_id,
            "title": f"{title}(档案)",
            "key": f"product_archive_{index:02d}",
            "primaryKey": False,
            "show": True,
            "width": 180,
            "align": "left",
        }
        if title in PRODUCT_DATE_HEADERS:
            column.update(type="date", dateFormat="YYYY-MM-DD HH:mm:ss")
        elif title in PRODUCT_NUMBER_HEADERS:
            number_format = "1,000.00" if title in product_amount_headers else "1000"
            column.update(type="number", numberFormat=number_format, align="right")
        else:
            column["type"] = "text"
        columns.append(column)
        next_id += 1
    return columns


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def normalize_number(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ImportFailure(f"布尔值不能作为数字导入: {value}")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ImportFailure(f"无法转换数字: {value}") from exc
    if not math.isfinite(number):
        raise ImportFailure(f"数字不是有限值: {value}")
    if number.is_integer():
        return int(number)
    return number


def normalize_date(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime_time.min)
    else:
        text = str(value).strip()
        if not text:
            return None
        parsed = None
        for date_format in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(text, date_format)
                break
            except ValueError:
                continue
        if parsed is None:
            raise ImportFailure(f"无法转换日期: {value}")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=SHANGHAI)
    return int(parsed.timestamp() * 1000)


def normalize_sequence(value: Any) -> int:
    number = normalize_number(value)
    if number is None or isinstance(number, float):
        raise ImportFailure(f"序号必须是整数: {value}")
    return number


def build_record(
    row: list[Any] | tuple[Any, ...],
    *,
    field_ids: dict[str, int],
    schema_id: int,
    schema_version: int,
    app_id: int,
    import_batch: str,
    product: dict[str, Any] | None = None,
) -> dict[str, Any]:
    sequence = normalize_sequence(row[0])
    import_key = f"{import_batch}:{sequence}"
    data: dict[str, Any] = {str(field_ids[IMPORT_KEY_TITLE]): import_key}

    for index, title in enumerate(SOURCE_HEADERS):
        value = row[index] if index < len(row) else None
        if title in DATE_HEADERS:
            normalized = normalize_date(value)
        elif title in NUMBER_HEADERS:
            normalized = normalize_number(value)
        else:
            normalized = normalize_text(value)
        if normalized is not None:
            data[str(field_ids[title])] = normalized

    if product is not None:
        for title in PRODUCT_HEADERS:
            if title in PRODUCT_EXCLUDED_HEADERS:
                continue
            value = product.get(title)
            if title in PRODUCT_DATE_HEADERS:
                normalized = normalize_date(value)
            elif title in PRODUCT_NUMBER_HEADERS:
                normalized = normalize_number(value)
            else:
                normalized = normalize_text(value)
            if normalized is not None:
                data[str(field_ids[f"{title}(档案)"])] = normalized

    return {
        "schemaId": schema_id,
        "version": schema_version,
        "appId": app_id,
        "primaryValue": import_key,
        "data": data,
    }


def source_sort_key(path: Path) -> tuple[int, str]:
    stem = path.stem
    if stem.endswith("-1"):
        return (1, path.name)
    if stem.endswith("-2"):
        return (2, path.name)
    return (0, path.name)


def source_workbooks(source_dir: Path) -> list[Path]:
    paths = sorted(source_dir.glob("*.xlsx"), key=source_sort_key)
    if not paths:
        raise ImportFailure(f"目录中没有 xlsx 文件: {source_dir}")
    return paths


def iter_source_rows(source_dir: Path) -> Iterator[tuple[Any, ...]]:
    for path in source_workbooks(source_dir):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            header = tuple(
                cell.value
                for cell in next(
                    sheet.iter_rows(
                        min_row=2,
                        max_row=2,
                        max_col=len(SOURCE_HEADERS),
                    )
                )
            )
            if list(header) != SOURCE_HEADERS:
                raise ImportFailure(f"文件表头不一致: {path.name}")
            for row in sheet.iter_rows(
                min_row=3,
                max_col=len(SOURCE_HEADERS),
                values_only=True,
            ):
                if normalize_text(row[0]) == "汇总":
                    continue
                if all(value is None for value in row):
                    continue
                yield row
        finally:
            workbook.close()


def collect_order_product_codes(source_dir: Path) -> set[str]:
    return {normalize_text(row[5]) for row in iter_source_rows(source_dir)}


def load_product_lookup(
    product_file: Path, required_codes: set[str]
) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(product_file, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [normalize_text(value) for value in next(rows)]
        except StopIteration as exc:
            raise ImportFailure(f"商品档案为空: {product_file}") from exc
        if headers != PRODUCT_HEADERS:
            raise ImportFailure(f"商品档案表头不一致: {product_file.name}")

        lookup: dict[str, dict[str, Any]] = {}
        for row in rows:
            sku = normalize_text(row[1] if len(row) > 1 else None)
            if not sku or sku not in required_codes:
                continue
            if sku in lookup:
                raise ImportFailure(f"商品档案 SKU聚合键重复: {sku}")
            lookup[sku] = {
                title: row[index] if index < len(row) else None
                for index, title in enumerate(PRODUCT_HEADERS)
            }
    finally:
        workbook.close()

    missing = sorted(required_codes - lookup.keys())
    if missing:
        preview = ",".join(missing[:20])
        raise ImportFailure(
            f"商品档案缺少 {len(missing)} 个订单 SKU，示例: {preview}"
        )
    return lookup


def source_fingerprint(source_dir: Path, product_file: Path | None = None) -> str:
    files = [
        {"name": path.name, "size": path.stat().st_size, "mtime_ns": path.stat().st_mtime_ns}
        for path in source_workbooks(source_dir)
    ]
    if product_file is not None:
        files.append(
            {
                "name": product_file.name,
                "size": product_file.stat().st_size,
                "mtime_ns": product_file.stat().st_mtime_ns,
            }
        )
    raw = json.dumps(files, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def run_cli_json(arguments: list[str]) -> dict[str, Any]:
    command = ["lingtong-cli", *arguments, "--envelope", "--format", "json"]
    completed = subprocess.run(command, capture_output=True, text=True, check=False)
    if completed.returncode != 0:
        raise ImportFailure(completed.stderr.strip() or completed.stdout.strip())
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise ImportFailure(f"CLI 返回了非 JSON 数据: {completed.stdout[:500]}") from exc
    if payload.get("ok") is not True:
        raise ImportFailure(str(payload.get("error") or payload))
    return payload


def find_schema(value: Any) -> dict[str, Any] | None:
    if isinstance(value, dict):
        if isinstance(value.get("columnsSchema"), list):
            return value
        for nested in value.values():
            found = find_schema(nested)
            if found is not None:
                return found
    elif isinstance(value, list):
        for nested in value:
            found = find_schema(nested)
            if found is not None:
                return found
    return None


def load_table_schema(table_id: int) -> tuple[int, int, dict[str, int]]:
    payload = run_cli_json(
        ["table", "schema", "query", "--basic-data-id", str(table_id)]
    )
    schema = find_schema(payload)
    if schema is None:
        raise ImportFailure("无法从 CLI 响应中找到表格 Schema")
    field_ids = {
        column["title"]: int(column["id"])
        for column in schema["columnsSchema"]
        if column.get("title") and column.get("id") is not None
    }
    expected_titles = [
        IMPORT_KEY_TITLE,
        *SOURCE_HEADERS,
        *(
            f"{title}(档案)"
            for title in PRODUCT_HEADERS
            if title not in PRODUCT_EXCLUDED_HEADERS
        ),
    ]
    missing = [title for title in expected_titles if title not in field_ids]
    if missing:
        raise ImportFailure(f"表格缺少字段: {','.join(missing)}")
    return int(schema["id"]), int(schema["version"]), field_ids


def load_table_count(table_id: int, schema_id: int, schema_version: int) -> int:
    payload = run_cli_json(
        [
            "table",
            "data",
            "count",
            "--schema-id",
            str(schema_id),
            "--version",
            str(schema_version),
            "--basic-data-id",
            str(table_id),
        ]
    )
    try:
        return int(payload["data"]["result"])
    except (KeyError, TypeError, ValueError) as exc:
        raise ImportFailure("无法从 CLI 响应中读取表格记录数") from exc


def load_last_import_key(table_id: int, remote_count: int) -> str:
    if remote_count <= 0:
        return ""
    page_size = 1000
    page = (remote_count - 1) // page_size + 1
    payload = run_cli_json(
        [
            "table",
            "data",
            "query",
            "--basic-data-id",
            str(table_id),
            "--page",
            str(page),
            "--page-size",
            str(page_size),
        ]
    )
    try:
        records = payload["data"]["result"]["data"]
        return normalize_text(records[-1]["data"]["1"])
    except (IndexError, KeyError, TypeError) as exc:
        raise ImportFailure("无法读取远端最后一条导入追踪键") from exc


def reconcile_remote_progress(
    *,
    processed: int,
    last_sequence: int,
    remote_count: int,
    last_import_key: str,
    import_batch: str,
) -> tuple[int, int]:
    if remote_count < processed:
        raise ImportFailure(
            f"远端记录数 {remote_count} 小于本地断点 {processed}，停止导入"
        )
    if remote_count == processed:
        return processed, last_sequence

    prefix = f"{import_batch}:"
    if not last_import_key.startswith(prefix):
        raise ImportFailure(f"远端最后一条追踪键不属于当前导入批次: {last_import_key}")
    try:
        remote_sequence = int(last_import_key[len(prefix) :])
    except ValueError as exc:
        raise ImportFailure(f"远端最后一条追踪键格式错误: {last_import_key}") from exc
    if remote_sequence <= last_sequence:
        raise ImportFailure(
            f"远端追踪键未领先本地断点: {last_import_key}, 本地序号 {last_sequence}"
        )
    return remote_count, remote_sequence


def read_cli_connection() -> tuple[str, str]:
    token = os.environ.get("LINGTONG_TOKEN")
    host_override = os.environ.get("LINGTONG_HOST")
    config_path = Path.home() / ".lingtong-cli" / "config.yaml"
    host = "https://app1.ltpass.com"
    current_auth = ""
    if config_path.exists():
        for raw_line in config_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if line.startswith("host:"):
                host = line.split(":", 1)[1].strip()
            elif line.startswith("currentAuth:"):
                current_auth = line.split(":", 1)[1].strip()
    if host_override:
        host = host_override
    if token:
        return host.rstrip("/"), token

    account = f"auth:{current_auth}" if current_auth else "default"
    completed = subprocess.run(
        [
            "security",
            "find-generic-password",
            "-s",
            "lingtong-cli",
            "-a",
            account,
            "-w",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if completed.returncode != 0:
        raise ImportFailure("无法从 macOS Keychain 读取 lingtong-cli Token")
    encoded = completed.stdout.strip()
    if encoded.startswith(KEYRING_BASE64_PREFIX):
        token = base64.b64decode(encoded[len(KEYRING_BASE64_PREFIX) :]).decode("utf-8")
    elif encoded.startswith(KEYRING_HEX_PREFIX):
        token = bytes.fromhex(encoded[len(KEYRING_HEX_PREFIX) :]).decode("utf-8")
    else:
        token = encoded
    return host.rstrip("/"), token


def validate_workflow_response(payload: dict[str, Any]) -> None:
    def visit(value: Any) -> None:
        if isinstance(value, dict):
            if value.get("success") is False:
                raise ImportFailure(str(value.get("msg") or value.get("message") or value))
            code = value.get("code")
            if isinstance(code, int) and code not in (0, 200, 10000):
                raise ImportFailure(str(value.get("msg") or value.get("message") or value))
            body = value.get("body")
            if isinstance(body, str) and body.lstrip().startswith(("{", "[")):
                try:
                    visit(json.loads(body))
                except json.JSONDecodeError:
                    pass
            for nested in value.values():
                if nested is not body:
                    visit(nested)
        elif isinstance(value, list):
            for nested in value:
                visit(nested)

    visit(payload)


def classify_failed_batch_count(
    processed: int, batch_size: int, remote_count: int
) -> str:
    if remote_count == processed:
        return "retry"
    if remote_count == processed + batch_size:
        return "committed"
    raise ImportFailure(
        f"批次可能部分提交：本地断点 {processed}，批次 {batch_size}，远端 {remote_count}"
    )


def build_workflow_request(
    *, account_name: str, table_id: int, records: list[dict[str, Any]]
) -> dict[str, Any]:
    return {
        "connector": "ltTable",
        "method": "LtTableRecordSave",
        "authAccount": account_name,
        "env": "prod",
        "body": {"basicDataId": table_id, "recordList": records},
    }


def post_workflow_batch(
    *,
    host: str,
    token: str,
    app_tag: str,
    account_name: str,
    table_id: int,
    records: list[dict[str, Any]],
    timeout: int,
) -> dict[str, Any]:
    request_payload = build_workflow_request(
        account_name=account_name, table_id=table_id, records=records
    )
    request = urllib.request.Request(
        f"{host}/gw/{app_tag}/workflows/run",
        data=json.dumps(request_payload, ensure_ascii=False, separators=(",", ":")).encode(
            "utf-8"
        ),
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")
        raise ImportFailure(f"工作流 HTTP {exc.code}: {detail[:1000]}") from exc
    except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as exc:
        raise ImportFailure(f"工作流请求失败: {exc}") from exc
    validate_workflow_response(payload)
    return payload


def load_checkpoint(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ImportFailure(f"断点文件损坏: {path}") from exc


def save_checkpoint(path: Path, state: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    temporary.replace(path)


def import_rows(args: argparse.Namespace) -> int:
    source_dir = args.source_dir.expanduser().resolve()
    product_file = args.product_file.expanduser().resolve()
    if not product_file.is_file():
        raise ImportFailure(f"商品档案不存在: {product_file}")
    checkpoint_path = (
        args.checkpoint.expanduser().resolve()
        if args.checkpoint
        else source_dir / f".lingtong-import-{args.table_id}.json"
    )
    fingerprint = source_fingerprint(source_dir, product_file)
    required_codes = collect_order_product_codes(source_dir)
    product_lookup = load_product_lookup(product_file, required_codes)
    schema_id, schema_version, field_ids = load_table_schema(args.table_id)
    checkpoint = None if args.restart else load_checkpoint(checkpoint_path)
    if checkpoint:
        if checkpoint.get("tableId") != args.table_id:
            raise ImportFailure("断点文件中的 tableId 与当前参数不一致")
        if checkpoint.get("sourceFingerprint") != fingerprint:
            raise ImportFailure("源文件已变化，不能使用现有断点继续导入")
    last_sequence = int(checkpoint.get("lastSequence", 0)) if checkpoint else 0
    processed = int(checkpoint.get("processed", 0)) if checkpoint else 0

    host = token = ""
    workflow_token = ""
    if not args.dry_run:
        host, token = read_cli_connection()
        workflow_token = os.environ.get("LINGTONG_WORKFLOW_SECRET", token)
        remote_count = load_table_count(args.table_id, schema_id, schema_version)
        if remote_count != processed:
            last_import_key = load_last_import_key(args.table_id, remote_count)
            processed, last_sequence = reconcile_remote_progress(
                processed=processed,
                last_sequence=last_sequence,
                remote_count=remote_count,
                last_import_key=last_import_key,
                import_batch=args.import_batch,
            )
            save_checkpoint(
                checkpoint_path,
                {
                    "tableId": args.table_id,
                    "schemaId": schema_id,
                    "schemaVersion": schema_version,
                    "sourceFingerprint": fingerprint,
                    "importBatch": args.import_batch,
                    "lastSequence": last_sequence,
                    "processed": processed,
                    "updatedAt": datetime.now(tz=SHANGHAI).isoformat(),
                },
            )
            print(
                f"远端领先本地断点，已按追踪键安全恢复到 {processed:,} 条，"
                f"序号 {last_sequence:,}",
                file=sys.stderr,
            )

    batch: list[dict[str, Any]] = []
    batch_last_sequence = last_sequence
    imported_this_run = 0
    started = time.monotonic()

    def flush() -> None:
        nonlocal batch, processed, imported_this_run
        if not batch:
            return
        if not args.dry_run:
            last_error = None
            for attempt in range(1, args.max_retries + 1):
                try:
                    post_workflow_batch(
                        host=host,
                        token=workflow_token,
                        app_tag=args.app_tag,
                        account_name=args.account_name,
                        table_id=args.table_id,
                        records=batch,
                        timeout=args.timeout,
                    )
                    last_error = None
                    break
                except ImportFailure as exc:
                    try:
                        remote_count = load_table_count(
                            args.table_id, schema_id, schema_version
                        )
                        count_state = classify_failed_batch_count(
                            processed, len(batch), remote_count
                        )
                    except ImportFailure as count_error:
                        raise ImportFailure(
                            f"批次失败且无法安全判断提交状态: {exc}; {count_error}"
                        ) from count_error
                    if count_state == "committed":
                        print(
                            "批次响应失败，但远端计数确认已完整提交，继续推进断点",
                            file=sys.stderr,
                        )
                        last_error = None
                        break
                    last_error = exc
                    if attempt == args.max_retries:
                        break
                    delay = min(30, 2 ** (attempt - 1))
                    print(
                        f"批次失败，{delay} 秒后进行第 {attempt + 1} 次尝试: {exc}",
                        file=sys.stderr,
                    )
                    time.sleep(delay)
            if last_error is not None:
                raise last_error

        processed += len(batch)
        imported_this_run += len(batch)
        state = {
            "tableId": args.table_id,
            "schemaId": schema_id,
            "schemaVersion": schema_version,
            "sourceFingerprint": fingerprint,
            "importBatch": args.import_batch,
            "lastSequence": batch_last_sequence,
            "processed": processed,
            "updatedAt": datetime.now(tz=SHANGHAI).isoformat(),
        }
        if not args.dry_run:
            save_checkpoint(checkpoint_path, state)
        elapsed = max(time.monotonic() - started, 0.001)
        print(
            f"已处理 {processed:,} 条，本次 {imported_this_run:,} 条，"
            f"最新序号 {batch_last_sequence:,}，平均 {imported_this_run / elapsed:.1f} 条/秒",
            flush=True,
        )
        batch = []

    for row in iter_source_rows(source_dir):
        sequence = normalize_sequence(row[0])
        if sequence <= last_sequence:
            continue
        batch.append(
            build_record(
                row,
                field_ids=field_ids,
                schema_id=schema_id,
                schema_version=schema_version,
                app_id=args.app_id,
                import_batch=args.import_batch,
                product=product_lookup[normalize_text(row[5])],
            )
        )
        batch_last_sequence = sequence
        if len(batch) >= args.batch_size:
            flush()
        if args.limit and imported_this_run + len(batch) >= args.limit:
            remaining = args.limit - imported_this_run
            if remaining < len(batch):
                batch = batch[:remaining]
                batch_last_sequence = int(batch[-1]["primaryValue"].rsplit(":", 1)[1])
            flush()
            break
    else:
        flush()

    print(
        json.dumps(
            {
                "dryRun": args.dry_run,
                "tableId": args.table_id,
                "schemaId": schema_id,
                "processed": processed,
                "importedThisRun": imported_this_run,
                "lastSequence": batch_last_sequence,
                "checkpoint": str(checkpoint_path),
            },
            ensure_ascii=False,
        )
    )
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--product-file", type=Path, required=True)
    parser.add_argument("--table-id", type=int, required=True)
    parser.add_argument("--app-tag", required=True)
    parser.add_argument("--account-name", default="纤莉秀表格导入")
    parser.add_argument("--app-id", type=int, default=213)
    parser.add_argument("--import-batch", default="20260914222134")
    parser.add_argument("--batch-size", type=int, default=500)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--max-retries", type=int, default=5)
    parser.add_argument("--timeout", type=int, default=180)
    parser.add_argument("--checkpoint", type=Path)
    parser.add_argument("--restart", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if args.batch_size <= 0:
        parser.error("--batch-size 必须大于 0")
    if args.limit < 0:
        parser.error("--limit 不能小于 0")
    return args


def main() -> int:
    try:
        return import_rows(parse_args())
    except ImportFailure as exc:
        print(f"导入失败: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
